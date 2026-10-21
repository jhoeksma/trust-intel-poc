from app.sources.loaders import load_generic, load_rtt


def csv(*lines):
    return ("\n".join(lines) + "\n").encode()


def test_long_format_nof():
    data = csv("Org Code,Org Name,Period,Metric ID,Metric Name,Metric Value",
               "RXX,Alpha Trust,Q1 2026/27,OF5200,Segment,2",
               "RXX,Alpha Trust,Q4 2025/26,OF5200,Segment,3",
               "RYY,Beta Trust,Q1 2026/27,OF5200,Segment,4")
    r = load_generic("nof", "trust-data.csv", data)
    assert r.profile["shape"] == "long"
    assert len(r.metrics) == 3 and set(r.orgs) == {"RXX", "RYY"}
    assert {m["period"] for m in r.metrics} == {"Q1 2026/27", "Q4 2025/26"}
    assert r.metrics[0]["metric_code"] == "nof:of5200"


def test_wide_format_with_suppression_markers():
    data = csv("Trust Code,Trust Name,Energy cost,Backlog maintenance,Notes",
               "RXX,Alpha,\"1,200\",*,hello",
               "RYY,Beta,300,45,there",
               "RZZ,Gamma,400,50,again")
    r = load_generic("eric", "ERIC - 2024_25 - Trust data.csv", data)
    codes = {m["metric_code"] for m in r.metrics}
    assert "eric:energy_cost" in codes and "eric:notes" not in codes
    assert all(m["period"] == "2024/25" for m in r.metrics)
    assert [m["value"] for m in r.metrics if m["ods_code"] == "RXX"] == [1200.0]  # '*' dropped


def test_no_org_column_is_profile_only():
    r = load_generic("x", "f.csv", csv("a,b", "1,2", "3,4"))
    assert r.metrics == [] and r.mapper_status == "needs_review"


def test_rtt_aggregation():
    hdr = ("Period,Provider Org Code,Provider Org Name,Commissioner Org Code,RTT Part Type,Treatment Function Code,"
           "Gt 00 To 01 Weeks SUM 1,Gt 17 To 18 Weeks SUM 1,Gt 18 To 19 Weeks SUM 1,Gt 52 Weeks SUM 1,Total All")
    data = csv(hdr,
               "RTT-April-2026,RXX,Alpha,AAA,Part_2,C_999,50,10,20,20,100",
               "RTT-April-2026,RXX,Alpha,BBB,Part_2,C_999,30,10,0,10,50",
               "RTT-April-2026,RXX,Alpha,BBB,Part_1A,C_999,999,0,0,0,999",
               "RTT-April-2026,RXX,Alpha,BBB,Part_2,C_100,999,0,0,0,999")
    r = load_rtt("rtt", "rtt.csv", data)
    got = {m["metric_code"]: m["value"] for m in r.metrics}
    assert got["rtt:incomplete_pathways_total"] == 150
    assert got["rtt:incomplete_over_18_weeks"] == 50
    assert got["rtt:pct_within_18_weeks"] == round(100 * 100 / 150, 2)
    assert r.mapper_status == "ok"


def test_rtt_unexpected_columns_fall_back():
    r = load_rtt("rtt", "rtt.csv", csv("Org Code,Value", "RXX,1", "RYY,2", "RZZ,3"))
    assert r.mapper_status == "needs_review"


WF_HDR = "DATA_MONTH,ORG_CODE,ORG_NAME,MAIN_STAFF_GROUP,STAFF_GROUP_1,FTE"


def test_workforce_dimensions_are_summed_not_collapsed():
    data = csv(WF_HDR,
               "2026-07-31,RXX,Alpha,Nurses,A,10",
               "2026-07-31,RXX,Alpha,Nurses,B,5.5",
               "2026-07-31,RXX,Alpha,Doctors,C,4",
               "2026-07-31,RYY,Beta,Nurses,A,7")
    r = load_generic("workforce", "Core 8. NWD Staff Group - org, Jul-26.csv", data)
    got = {(m["ods_code"], m["metric_code"]): m["value"] for m in r.metrics}
    assert got[("RXX", "workforce:core_8:fte")] == 19.5
    assert got[("RXX", "workforce:core_8:fte:nurses")] == 15.5
    assert got[("RXX", "workforce:core_8:fte:doctors")] == 4
    assert got[("RYY", "workforce:core_8:fte")] == 7
    keys = [(m["ods_code"], m["metric_code"], m["period"]) for m in r.metrics]
    assert len(keys) == len(set(keys))  # one value per org, metric, period
    assert r.profile["breakdown_column"] == "MAIN_STAFF_GROUP"


def test_url_encoded_file_name_still_gives_period():
    data = csv("Trust Code,Trust Name,Energy cost", "RXX,Alpha,5", "RYY,Beta,6", "RZZ,Gamma,7")
    r = load_generic("eric", "ERIC%20-%202024_25%20-%20Trust%20data.csv", data)
    assert {m["period"] for m in r.metrics} == {"2024/25"}


def test_long_file_counts_suppressed_values():
    data = csv("ORG_CODE,ORG_DESCRIPTION,REPORTING_PERIOD,MEASURE_TYPE,MEASURE,MEASURE_VALUE",
               "RXX,Alpha,2425,Summary,FAE_SUM,100",
               "RYY,Beta,2425,Summary,FAE_SUM,*",
               "RZZ,Gamma,2425,Summary,FAE_SUM,50")
    r = load_generic("hes", "hosp-epis-stat-admi-pla-2024-25.csv", data)
    assert r.profile["suppressed_values"] == 1 and len(r.metrics) == 2
    assert "[Summary]" in r.metrics[0]["metric_label"]


# ---- NOF (trust data and league table)
NOF_TD = ("Quarter,Region,Trust_type,Trust_subtype,Trust_code,Trust_name,Domain,Sub-domain,Metric_ID,Metric_description,"
          "Reporting_date,Units,Value,Median_value,Lower_quartile,Upper_quartile,Rank\n"
          "Q1 2026/27,London,Acute trust,Acute - Large,RXX,Alpha NHS Trust,Summary,Summary,OF5200,Adjusted segment,Q1 26/27,segment,3,2,2,3,40\n"
          "Q1 2026/27,London,Acute trust,Acute - Large,RXX,Alpha NHS Trust,Summary,Summary,OF5202,Average metric score,Q1 26/27,score,2.7000000000000002,2.5,2,3,\n"
          "Q1 2026/27,London,Acute trust,Acute - Large,RXX,Alpha NHS Trust,People,People,OF0084,Staff\xa0engagement,2025,out of 10,6.8,7,6,7.5,12\n")
NOF_LT = ("Quarter,Region,Trust_type,Trust_subtype,Trust_code,Trust_name,Average_score,Likely_range_of_average_score,Segment,"
          "Trust_in_financial_deficit,Rank,Likely_range_of_rank\n"
          "Q1 26/27,London,Acute trust,Acute - Large,RXX,Alpha NHS Trust,2.7,2.4 to 3.1,3,Yes,40,31 to 52\n")


def test_nof_trust_data():
    from app.sources.nof import load_nof
    r = load_nof("nof", "NOF - Q1 2026_27 - Trust data - Acute.csv.gz", NOF_TD.encode())
    by = {m["metric_code"]: m for m in r.metrics}
    assert by["nof:of5200"]["value"] == 3 and by["nof:of5200"]["period"] == "Q1 2026/27"
    assert by["nof:of5200:rank"]["value"] == 40
    assert by["nof:of5202"]["value"] == 2.7  # float noise removed
    assert "nof:of5202:rank" not in by  # blank rank is not a zero
    assert by["nof:of0084"]["unit"] == "score10" and "Staff engagement" in by["nof:of0084"]["metric_label"]
    assert r.orgs["RXX"] == ("Alpha NHS Trust", "Acute trust")


def test_nof_league_table_loads_only_new_fields():
    from app.sources.nof import load_nof
    r = load_nof("nof", "league.csv", NOF_LT.encode())
    by = {m["metric_code"]: m["value"] for m in r.metrics}
    assert by == {"nof:league_rank": 40, "nof:league_score_low": 2.4, "nof:league_score_high": 3.1,
                  "nof:league_rank_low": 31, "nof:league_rank_high": 52}
    assert {m["period"] for m in r.metrics} == {"Q1 2026/27"}  # "Q1 26/27" normalised


def test_nof_unrecognised_layout_is_profile_only():
    from app.sources.nof import load_nof
    r = load_nof("nof", "x.csv", b"a,b\n1,2\n")
    assert r.metrics == [] and r.mapper_status == "needs_review"


def test_strain_uses_adjusted_segment_not_domain_or_unadjusted():
    from app.indicators import strain_from_metrics
    ms = [{"source": "nof", "metric_code": "nof:of5203", "label": "Unadjusted segment", "value": 1, "period": "p"},
          {"source": "nof", "metric_code": "nof:of4300", "label": "domain segment", "value": 2, "period": "p"},
          {"source": "nof", "metric_code": "nof:of5200", "label": "Adjusted segment", "value": 4, "period": "p"}]
    assert strain_from_metrics(ms)["score"] == 4


# ---- RTT
def _rtt_csv(rows):
    head = ["Period", "Provider Org Code", "Provider Org Name", "RTT Part Type", "Treatment Function Code"]
    head += [f"Gt {i:02d} To {i + 1:02d} Weeks SUM 1" for i in range(0, 104)] + ["Gt 104 Weeks SUM 1", "Total", "Patients with unknown clock start date", "Total All"]
    out = [",".join(head)]
    for part, tfc, buckets in rows:
        vec = [0] * 105
        for wk, n in buckets.items():
            vec[wk] = n
        out.append(",".join(["RTT-August-2026", "RXX", "Alpha", part, tfc] + [str(v) for v in vec] + ["", "", str(sum(vec))]))
    return "\n".join(out).encode()


def test_rtt_incomplete_excludes_unknown_clock_and_other_specialties():
    from app.sources.loaders import load_rtt
    data = _rtt_csv([
        ("Part_2", "C_999", {0: 50, 17: 30, 18: 10, 60: 10}),  # 100 waiting: 80 within 18 weeks, 10 over 52
        ("Part_2", "C_100", {0: 999}),  # specialty row: ignored
        ("Part_2A", "C_999", {5: 500}),  # unknown clock start: must not enter the distribution
        ("Part_1A", "C_999", {3: 7}),
    ])
    r = load_rtt("rtt", "x.csv", data)
    by = {m["metric_code"]: m["value"] for m in r.metrics}
    assert by["rtt:incomplete_pathways_total"] == 100
    assert by["rtt:pct_within_18_weeks"] == 80.0  # bucket 'Gt 17 To 18' is within 18 weeks
    assert by["rtt:incomplete_over_52_weeks"] == 10 and by["rtt:pct_over_52_weeks"] == 10.0
    assert by["rtt:unknown_clock_start"] == 500 and by["rtt:completed_admitted"] == 7
    assert 0 <= by["rtt:incomplete_median_wait_weeks"] <= 1  # 50 of 100 sit in the first week bucket
    assert by["rtt:incomplete_p92_wait_weeks"] >= 18


def test_nof_emits_peer_group_once_per_trust():
    from app.sources.nof import load_nof, peer_group_code
    assert peer_group_code("Acute trust") == 1 and peer_group_code("Mental health and community trust") == 2
    assert peer_group_code("Ambulance trust") == 3 and peer_group_code("Independent") is None
    hdr = "Trust_code,Trust_name,Trust_type,Quarter,Metric_ID,Metric_description,Value,Units\n"
    body = ("RXX,Alpha,Mental health and community trust,Q1 2026/27,OF5200,Segment,2,segment\n"
            "RXX,Alpha,Mental health and community trust,Q1 2026/27,OF5202,Score,2.1,score\n")
    r = load_nof("nof", "NOF - Q1 2026_27 - Trust data - Mental health and community.csv", (hdr + body).encode())
    g = [m for m in r.metrics if m["metric_code"] == "nof:peer_group"]
    assert len(g) == 1 and g[0]["value"] == 3.0 or g[0]["value"] == 2.0
    assert g[0]["value"] == 2.0 and g[0]["unit"] == "group"


def test_nof_peer_subtype_codes():
    from app.sources.nof import load_nof, peer_subtype_code
    cases = {"Acute - Teaching": 11, "Acute - Large": 12, "Acute - Medium": 13, "Acute - Multi-Service": 14, "Acute - Small": 15,
             "Acute - Specialist": 16, "Community": 21, "Mental Health and Learning Disability": 22, "Care Trust": 23, "Ambulance": 31}
    for text, code in cases.items():
        assert peer_subtype_code(text) == code, text
    assert peer_subtype_code("") is None
    hdr = "Trust_code,Trust_name,Trust_type,Trust_subtype,Quarter,Metric_ID,Metric_description,Value,Units\n"
    body = ("RXX,Alpha,Acute trust,Acute - Teaching,Q1 2026/27,OF5200,Segment,2,segment\n"
            "RXX,Alpha,Acute trust,Acute - Teaching,Q1 2026/27,OF5202,Score,2.1,score\n")
    r = load_nof("nof", "NOF - Q1 2026_27 - Trust data - Acute.csv", (hdr + body).encode())
    s = [m for m in r.metrics if m["metric_code"] == "nof:peer_subtype"]
    assert len(s) == 1 and s[0]["value"] == 11.0


def _dma_workbook() -> bytes:
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Introduction "
    p = wb.create_sheet("SC - 2025 Pillar Summary")
    p.append([None, "2025 DMA"])
    p.append([None, "Region", "ICS", "Provider", "Care Setting", "Well Led", "Ensure Smart Foundations", "Safe Practice",
              "Support Workforce", "Empower People", "Improve Care", "Healthy Populations"])
    p.append([None, "R", "I", "ALPHA NHS TRUST", "Acute", 2.8, 2.4, 2, 2, 1.9, 2, 1.4])
    p.append([None, "R", "I", "ALPHA NHS TRUST", "Community", 3, 3, 3, 3, 3, 3, 3])
    p.append([None, "R", "I", "BETA & GAMMA NHS FOUNDATION TRUST", "Mental Health", 2, 2, 2, 2, 2, 2, 2])
    p.append([None, "R", "I", "NOBODY NHS TRUST", "Acute", 2, 2, 2, 2, 2, 2, 2])
    o = wb.create_sheet("SC - 24 vs 25 Overall DMA")
    o.append([None, "x"])
    o.append([None, "Region", "ICS", "Provider", "Care Setting", "2025 Average", "2024 Average"])
    o.append([None, "R", "I", "ALPHA NHS TRUST", "Acute", 2.1, 2.5])
    o.append([None, "R", "I", "ALPHA NHS TRUST", "Community", 3.2, 3.0])
    o.append([None, "R", "I", "BETA & GAMMA NHS FOUNDATION TRUST", "Mental Health", 2.0, None])
    o.append([None, "R", "I", "NOBODY NHS TRUST", "Acute", 2.2, 2.2])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_dma_primary_setting_pillars_and_unmatched():
    from app.sources.dma import load_dma
    known = {"RXX": "Alpha NHS Trust", "RYY": "Beta and Gamma NHS Foundation Trust"}
    r = load_dma("dma", "dma.xlsx", _dma_workbook(), known_orgs=known, overrides={})
    by = {(m["ods_code"], m["metric_code"], m["period"]): m for m in r.metrics}
    assert by[("RXX", "dma:overall", "2025")]["value"] == 2.1 and by[("RXX", "dma:overall", "2024")]["value"] == 2.5  # acute is primary
    assert by[("RXX", "dma:overall_community", "2025")]["value"] == 3.2  # other setting kept apart, not the headline
    assert by[("RXX", "dma:pillar_well_led", "2025")]["value"] == 2.8  # pillars from the primary setting only
    assert by[("RXX", "dma:care_setting", "2025")]["value"] == 1
    assert ("RYY", "dma:overall", "2024") not in by  # blank 2024 not stored as zero
    assert by[("RYY", "dma:care_setting", "2025")]["value"] == 3
    assert r.profile["unmatched_providers"] == ["NOBODY NHS TRUST"]
    assert set(r.orgs) == {"RXX", "RYY"}


def test_dma_override_must_point_at_a_known_org():
    from app.sources.dma import load_dma
    r = load_dma("dma", "dma.xlsx", _dma_workbook(), known_orgs={"RXX": "Alpha NHS Trust"},
                 overrides={"nobody": "ZZZ", "beta and gamma": "RXX"})
    assert "ZZZ" in " ".join(r.profile["unmatched_providers"])  # unknown override code is reported, not created
    assert "RXX" in {m["ods_code"] for m in r.metrics}

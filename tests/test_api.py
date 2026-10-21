import io

import pytest
from fastapi.testclient import TestClient

from app.models import Metric, Org, SourceFile, TrustSystem


@pytest.fixture()
def client(db):
    from app.main import app
    sf = SourceFile(source_key="nof", status="loaded")
    db.add(sf)
    db.add_all([Org(ods_code="RXX", name="Alpha Trust"), Org(ods_code="RYY", name="Beta Trust")])
    db.flush()

    def m(org, code, label, period, sort, val, src="nof"):
        return Metric(ods_code=org, source_key=src, metric_code=code, metric_label=label, period=period,
                      period_sort=sort, period_type="quarter", value=val, source_file_id=sf.id)
    db.add_all([
        m("RXX", "nof:of5200", "Adjusted segment", "Q1 2026/27", 202604, 2),
        m("RXX", "nof:of5200", "Adjusted segment", "Q4 2025/26", 202601, 3),
        m("RYY", "nof:of5200", "Adjusted segment", "Q1 2026/27", 202604, 4),
        m("RXX", "rtt:pct_within_18_weeks", "RTT within 18w", "2026-04", 202604, 61.5, "rtt"),
    ])
    db.commit()
    return TestClient(app)


def test_health(client):
    assert client.get("/health").json()["status"] == "ok"


def test_trust_list_has_strain(client):
    rows = {r["ods_code"]: r for r in client.get("/api/trusts").json()}
    assert rows["RXX"]["strain"] == 2 and rows["RYY"]["strain"] == 4
    assert rows["RXX"]["rtt_pct_18w"] == 61.5


def test_scorecard_latest_and_previous(client):
    card = client.get("/api/trusts/rxx").json()
    seg = next(x for x in card["metrics"] if x["metric_code"] == "nof:of5200")
    assert seg["value"] == 2 and seg["previous_value"] == 3
    assert card["strain"]["score"] == 2


def test_unknown_trust_404(client):
    assert client.get("/api/trusts/ZZZ").status_code == 404


def test_compare(client):
    c = client.get("/api/compare?codes=RXX,RYY").json()
    row = next(r for r in c["rows"] if r["metric_code"] == "nof:of5200")
    assert row["values"]["RXX"]["value"] == 2 and row["values"]["RYY"]["value"] == 4


def test_systems_import_and_join_rate(client):
    csv_data = "ods_code,category,supplier,product,contract_end\nRXX,EPR,Acme,Acme EPR,31/03/2028\nQQQ,EPR,Other,X,2027-01-01\n"
    r = client.post("/api/systems/import", files={"file": ("s.csv", io.BytesIO(csv_data.encode()), "text/csv")})
    body = r.json()
    assert body["imported"] == 2 and body["codes_not_in_public_data"] == ["QQQ"] and body["join_rate"] == 0.5
    card = client.get("/api/trusts/RXX").json()
    assert card["systems"][0]["supplier"] == "Acme" and card["systems"][0]["contract_end"] == "2028-03-31"
    assert next(t for t in client.get("/api/trusts").json() if t["ods_code"] == "RXX")["suppliers"] == ["Acme"]


def test_sources_endpoint(client):
    s = client.get("/api/sources").json()
    assert {x["key"] for x in s["registry"]} >= {"nof", "rtt", "eric", "hes", "workforce"}
    assert s["coverage"][0]["source"] in {"nof", "rtt"}


def test_auth_enforced_when_required(client, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "auth_required", True)
    monkeypatch.setattr(settings, "basic_auth_user", "u")
    monkeypatch.setattr(settings, "basic_auth_password", "p")
    assert client.get("/api/trusts").status_code == 401
    assert client.get("/api/trusts", auth=("u", "p")).status_code == 200
    assert client.get("/health").status_code == 200


def test_reference_seed_and_coverage(client, db):
    from app.reference import load_reference, seed_orgs
    ref = load_reference()
    assert len(ref) > 200 and all(r["ods_code"] for r in ref) and len({r["ods_code"] for r in ref}) == len(ref)
    assert seed_orgs(db) > 0
    db.add(Metric(ods_code="ZZZ", source_key="x", metric_code="x:a", period="2026-04", period_sort=202604, period_type="month", value=1))
    db.commit()
    cov = client.get("/api/sources").json()["reference_coverage"]
    assert cov["reference_orgs"] == len(ref) and "ZZZ" in cov["public_data_not_in_reference"]


def test_successor_code_history_shows_on_current_trust(client, db):
    from app.reference import load_successors
    assert load_successors().get("RBZ") == "RH8"
    db.add_all([Org(ods_code="RBZ", name="Northern Devon", current_code="RH8"), Org(ods_code="RH8", name="Royal Devon")])
    db.add(Metric(ods_code="RBZ", source_key="rtt", metric_code="rtt:pct_within_18_weeks", period="2021-12", period_sort=202112,
                  period_type="month", value=55.0))
    db.commit()
    card = client.get("/api/trusts/RH8").json()
    assert any(m["metric_code"] == "rtt:pct_within_18_weeks" for m in card["metrics"])


@pytest.fixture()
def bench(db):
    """Twelve acute-looking trusts with a lower-is-better and a higher-is-better metric."""
    from app.benchmark import _cache
    _cache.clear()
    from app.main import app
    sf = SourceFile(source_key="nof", status="loaded")
    db.add(sf)
    db.flush()
    for i in range(12):
        code = f"R{i:02d}"
        db.add(Org(ods_code=code, name=f"Trust {i}"))
        db.add(Metric(ods_code=code, source_key="nof", metric_code="nof:of5200", metric_label="Segment", period="Q1 2026/27",
                      period_sort=202604, period_type="quarter", value=float(1 + i % 4), unit="segment", source_file_id=sf.id))
        db.add(Metric(ods_code=code, source_key="rtt", metric_code="rtt:pct_within_18_weeks", metric_label="18w", period="2026-08",
                      period_sort=202608, period_type="month", value=50.0 + i, unit="pct", source_file_id=sf.id))
        db.add(Metric(ods_code=code, source_key="rtt", metric_code="rtt:incomplete_pathways_total", metric_label="Waiting", period="2026-08",
                      period_sort=202608, period_type="month", value=1000.0 * (i + 1), unit="count", source_file_id=sf.id))
    db.add(Metric(ods_code="R00", source_key="nof", metric_code="nof:of5200:rank", metric_label="r", period="Q1 2026/27",
                  period_sort=202604, period_type="quarter", value=1.0, source_file_id=sf.id))
    db.commit()
    return TestClient(app)


def test_catalogue_excludes_companions_and_marks_polarity(bench):
    rows = {r["metric_code"]: r for r in bench.get("/api/benchmark/metrics").json()}
    assert "nof:of5200:rank" not in rows
    assert rows["nof:of5200"]["polarity"] == "lower" and rows["nof:of5200"]["core"]
    assert rows["rtt:pct_within_18_weeks"]["polarity"] == "higher"
    assert rows["rtt:incomplete_pathways_total"]["polarity"] is None


def test_metric_view_ranks_best_first(bench):
    v = bench.get("/api/benchmark/metric?code=rtt:pct_within_18_weeks").json()
    assert v["points"][0]["ods_code"] == "R11" and v["points"][-1]["ods_code"] == "R00"
    assert v["stats"]["n"] == 12 and v["stats"]["median"] == 55.5 and v["ranking"] == "best first"
    low = bench.get("/api/benchmark/metric?code=nof:of5200").json()
    assert low["points"][0]["value"] == 1.0
    neutral = bench.get("/api/benchmark/metric?code=rtt:incomplete_pathways_total").json()
    assert neutral["ranking"] == "highest first" and neutral["points"][0]["value"] == 12000


def test_trust_position_and_ties(bench):
    b = bench.get("/api/benchmark/trust/r11?codes=rtt:pct_within_18_weeks,nof:of5200&scope=national").json()
    rtt = next(m for m in b["metrics"] if m["metric_code"] == "rtt:pct_within_18_weeks")
    assert rtt["rank"] == 1 and rtt["n"] == 12 and rtt["better_than_pct"] == 100.0 and rtt["quartile"] == 1
    seg = next(m for m in b["metrics"] if m["metric_code"] == "nof:of5200")
    assert seg["value"] == 4.0 and seg["rank"] == 10 and seg["quartile"] == 4  # three trusts at 4: tied worst, rank 10
    assert b["peer_group"] == "acute"


def test_neutral_metric_has_no_quartile_and_peers_fall_back(bench):
    b = bench.get("/api/benchmark/trust/R03?codes=rtt:incomplete_pathways_total&scope=peers").json()
    m = b["metrics"][0]
    assert m["quartile"] is None and m["rank"] == 9 and m["scope_used"] == "peers"


def test_default_codes_are_core_and_bad_input(bench):
    b = bench.get("/api/benchmark/trust/R01").json()
    assert {m["metric_code"] for m in b["metrics"]} >= {"nof:of5200", "rtt:pct_within_18_weeks"}
    assert bench.get("/api/benchmark/trust/ZZZ").status_code == 404
    assert bench.get("/api/benchmark/trust/R01?scope=x").status_code == 422
    assert bench.get("/api/benchmark/metric?code=nope").status_code == 404
    assert bench.get("/api/benchmark/metric?code=nof:of5200&group=bogus").status_code == 422


def test_publisher_trust_type_overrides_inferred_group(bench, db):
    from app.benchmark import _cache
    sf = db.query(SourceFile).first()
    db.add(Metric(ods_code="R03", source_key="nof", metric_code="nof:peer_group", metric_label="type", period="Q1 2026/27",
                  period_sort=202604, period_type="quarter", value=2.0, unit="group", source_file_id=sf.id))
    db.commit()
    _cache.clear()
    assert bench.get("/api/benchmark/trust/R03").json()["peer_group"] == "mhc"
    assert bench.get("/api/benchmark/trust/R04").json()["peer_group"] == "acute"
    assert "nof:peer_group" not in {r["metric_code"] for r in bench.get("/api/benchmark/metrics").json()}


def test_merged_code_points_to_successor(bench, db):
    db.add_all([Org(ods_code="RNK", name="Tavistock"), Org(ods_code="G6V2S", name="North London NHS Foundation Trust")])
    db.commit()
    b = bench.get("/api/benchmark/trust/RNK").json()
    assert b["merged_into"] == "G6V2S" and b["merged_into_name"].startswith("North London")
    assert bench.get("/api/benchmark/trust/R01").json()["merged_into"] is None


@pytest.fixture()
def subs(bench, db):
    """R00-R05 Acute - Teaching (six, enough for peers), R06-R08 Acute - Small (three, too few), rest unclassified."""
    from app.benchmark import _cache
    sf = db.query(SourceFile).first()
    for i in range(12):
        code = 11 if i < 6 else 15 if i < 9 else None
        if code:
            db.add(Metric(ods_code=f"R{i:02d}", source_key="nof", metric_code="nof:peer_subtype", metric_label="sub", period="Q1 2026/27",
                          period_sort=202604, period_type="quarter", value=float(code), unit="group", source_file_id=sf.id))
    db.commit()
    _cache.clear()
    return bench


def test_group_options_and_subtype_filter(subs):
    o = subs.get("/api/benchmark/groups").json()
    assert {s["key"]: s["n"] for s in o["subtypes"]} == {"acute_teaching": 6, "acute_small": 3}
    assert {g["key"]: g["n"] for g in o["groups"]}["acute"] == 12
    v = subs.get("/api/benchmark/metric?code=rtt:pct_within_18_weeks&subtype=acute_teaching").json()
    assert v["stats"]["n"] == 6 and {p["subtype"] for p in v["points"]} == {"acute_teaching"}
    assert subs.get("/api/benchmark/metric?code=nof:of5200&subtype=nope").status_code == 422
    assert "nof:peer_subtype" not in {r["metric_code"] for r in subs.get("/api/benchmark/metrics").json()}


def test_subtype_scope_and_fallback(subs):
    t = subs.get("/api/benchmark/trust/R02?codes=rtt:pct_within_18_weeks&scope=subtype").json()
    m = t["metrics"][0]
    assert t["subtype_label"] == "Acute - Teaching" and m["scope_used"] == "subtype" and m["n"] == 6 and m["rank"] == 4
    small = subs.get("/api/benchmark/trust/R07?codes=rtt:pct_within_18_weeks&scope=subtype").json()["metrics"][0]
    assert small["scope_used"] == "peers" and small["n"] == 12  # only three small acute trusts: falls back to type
    plain = subs.get("/api/benchmark/trust/R10?codes=rtt:pct_within_18_weeks&scope=subtype").json()["metrics"][0]
    assert plain["scope_used"] == "peers"  # no subtype at all
    assert subs.get("/api/benchmark/trust/R02?scope=national").status_code == 200


def test_trust_list_carries_type_and_subtype(subs):
    rows = {r["ods_code"]: r for r in subs.get("/api/trusts").json()}
    assert rows["R01"]["subtype"] == "acute_teaching" and rows["R01"]["peer_group"] == "acute" and rows["R10"]["subtype"] is None


def test_trend_series_against_group(bench, db):
    from app.benchmark import _cache
    sf = db.query(SourceFile).first()
    for i in range(12):  # an earlier month: everyone five points worse
        db.add(Metric(ods_code=f"R{i:02d}", source_key="rtt", metric_code="rtt:pct_within_18_weeks", metric_label="18w", period="2026-07",
                      period_sort=202607, period_type="month", value=45.0 + i, unit="pct", source_file_id=sf.id))
    db.commit()
    _cache.clear()
    t = bench.get("/api/benchmark/trust/R11/trend?code=rtt:pct_within_18_weeks&scope=national").json()
    assert [p["period"] for p in t["points"]] == ["2026-07", "2026-08"] and t["periods"] == 2
    jul, aug = t["points"]
    assert jul["value"] == 56.0 and aug["value"] == 61.0 and jul["rank"] == 1 and aug["n"] == 12
    assert jul["median"] == 50.5 and aug["median"] == 55.5 and t["scope_used"] == "national"
    assert bench.get("/api/benchmark/trust/R11/trend?code=nope").status_code == 404
    assert bench.get("/api/benchmark/trust/R11/trend?code=rtt:pct_within_18_weeks&scope=x").status_code == 422
    cat = {r["metric_code"]: r for r in bench.get("/api/benchmark/metrics").json()}
    assert cat["rtt:pct_within_18_weeks"]["periods"] == 2 and cat["nof:of5200"]["periods"] == 1


def test_metric_trend_group_series(subs, db):
    from app.benchmark import _cache
    sf = db.query(SourceFile).first()
    for i in range(12):  # an earlier month; R05 did not report it
        if i != 5:
            db.add(Metric(ods_code=f"R{i:02d}", source_key="rtt", metric_code="rtt:pct_within_18_weeks", metric_label="18w", period="2026-07",
                          period_sort=202607, period_type="month", value=40.0 + i, unit="pct", source_file_id=sf.id))
    db.commit()
    _cache.clear()
    t = subs.get("/api/benchmark/metric/trend?code=rtt:pct_within_18_weeks&subtype=acute_teaching").json()
    assert [p["period"] for p in t["periods"]] == ["2026-07", "2026-08"]
    assert t["periods"][0]["n"] == 5 and t["periods"][1]["n"] == 6  # R05 missing in July
    assert len(t["series"]) == 6
    r05 = next(s for s in t["series"] if s["ods_code"] == "R05")
    assert r05["values"] == [None, 55.0]
    assert t["periods"][1]["median"] == 52.5
    assert subs.get("/api/benchmark/metric/trend?code=rtt:pct_within_18_weeks&months=1").json()["periods_shown"] == 1
    assert subs.get("/api/benchmark/metric/trend?code=nope").status_code == 404
    assert subs.get("/api/benchmark/metric/trend?code=rtt:pct_within_18_weeks&subtype=x").status_code == 422


@pytest.fixture()
def dma_client(db):
    """20 acute trusts: higher DMA goes with better (lower) NOF score and higher CQC rating; plus one ambulance trust."""
    from app.main import app
    sf = SourceFile(source_key="dma", status="loaded")
    db.add(sf)
    db.flush()
    rows = []

    def m(org, code, per, sort, val):
        rows.append(Metric(ods_code=org, source_key=code.split(":")[0], metric_code=code, metric_label=code, period=per,
                           period_sort=sort, period_type="year", value=val, source_file_id=sf.id))
    for i in range(20):
        o = f"T{i:02d}"
        db.add(Org(ods_code=o, name=f"Trust {i}"))
        m(o, "dma:overall", "2025", 202504, 1.5 + i * 0.1)
        m(o, "dma:overall", "2024", 202404, 1.4 + i * 0.1)
        m(o, "dma:pillar_well_led", "2025", 202504, 2 + (i % 5) * 0.1)
        m(o, "dma:care_setting", "2025", 202504, 1)
        m(o, "nof:peer_group", "Q1 2026/27", 202604, 1)
        m(o, "nof:of5202", "Q1 2026/27", 202604, 3.5 - i * 0.05)  # lower is better
        m(o, "nof:of5200", "Q1 2026/27", 202604, 4 if i < 5 else 3 if i < 10 else 2 if i < 15 else 1)
        m(o, "cqc:overall", "2026-10", 202610, 2 if i < 10 else 3)
    for i in range(20):  # first 8 on Alpha, next 8 on Beta, 2 on both, 2 with no EPR recorded
        sups = ["Alpha"] if i < 8 else ["Beta"] if i < 16 else ["Alpha", "Beta"] if i < 18 else []
        for sp in sups:
            db.add(TrustSystem(ods_code=f"T{i:02d}", category="EPR", supplier=sp))
        db.add(TrustSystem(ods_code=f"T{i:02d}", category="PACS", supplier="Pix"))
    db.add(Org(ods_code="AMB", name="Ambulance"))
    m("AMB", "dma:overall", "2025", 202504, 3.0)
    m("AMB", "nof:peer_group", "Q1 2026/27", 202604, 3)
    db.add_all(rows)
    db.commit()
    return TestClient(app)


def test_dma_correlation_signs_and_quartiles(dma_client):
    d = dma_client.get("/api/dma/correlation?group=acute").json()
    assert d["n_trusts"] == 20  # ambulance trust filtered out
    nof = d["correlations"]["overall"]["nof_score"]
    assert nof["rho"] > 0.95 and nof["lo"] <= nof["rho"] <= nof["hi"]  # positive means more mature does better, even though lower is better
    assert d["correlations"]["overall"]["cqc_overall"]["rho"] > 0.6
    assert d["correlations"]["overall"]["fin_domain"] is None  # no data, no invented correlation
    q = d["quartiles"]["nof_score"]
    assert [r["q"] for r in q] == [1, 2, 3, 4] and q[0]["mean"] > q[3]["mean"]  # least mature quartile has the worse (higher) score
    assert d["quartiles"]["cqc_overall"][3]["good_pct"] == 100.0 and d["quartiles"]["cqc_overall"][0]["good_pct"] == 0.0
    assert d["time"]["n"] == 20 and d["time"]["periods"] == [2024, 2025] and d["time"]["median_change"] == pytest.approx(0.1)
    assert d["time"]["improved"] == 20 and d["trusts"][0]["dma"]["overall_prev"] == pytest.approx(3.3)
    t = d["trusts"][0]
    assert t["ods_code"] == "T19" and t["dma"]["change"] == pytest.approx(0.1) and t["care_setting"] == "Acute"


def test_dma_correlation_all_and_bad_group(dma_client):
    assert dma_client.get("/api/dma/correlation").json()["n_trusts"] == 21
    assert dma_client.get("/api/dma/correlation?group=nope").status_code == 422


def test_systems_master_layout_maps_via_fhi_id_and_skips_junk(client, monkeypatch):
    from app import systems as sy
    monkeypatch.setattr(sy, "fhi_to_ods", lambda: {"7": "RXX", "8": "RYY"})
    hdr = "Trust_ID,Foundation_Trust_Name,Trust_Type,DepartmentName,SupplierName,SystemName,Has_System,Upd_Flag,Contract_expiration_notes\n"
    body = hdr + "7,A,Acute,EPR,Acme,Acme EPR,Yes,,31/03/2028\n7,A,Acute,PACS,Pixel,P,Yes,,\n8,B,Acute,EPR,Beta,B EPR,Yes,,\n8,B,Acute,EPR,Old,O,Yes,D,\n9,C,Acute,EPR,Zed,Z,Yes,,\n8,B,Acute,Pharmacy,,,No,,\n"
    r = client.post("/api/systems/import", files={"file": ("m.csv", io.BytesIO(body.encode()), "text/csv")}).json()
    assert r["format"] == "master" and r["rows_loaded"] == 3 and r["unmapped_organisations"] == 1
    assert r["skipped"]["deleted_flag"] == 1
    o = client.get("/api/systems/options").json()
    assert o["loaded"] and o["categories"][0] == {"category": "EPR", "trusts": 2}
    assert {x["supplier"] for x in o["suppliers"]["EPR"]} == {"Acme", "Beta"}


def test_dma_filter_by_epr_supplier(dma_client):
    base = dma_client.get("/api/dma/correlation?group=acute").json()
    assert base["systems_loaded"] and base["by_supplier_category"] == "EPR"
    rows = {r["supplier"]: r for r in base["by_supplier"]}
    assert rows["Alpha"]["n"] == 8 and rows["Beta"]["n"] == 8 and "Multiple suppliers" not in rows  # 2 multi-supplier trusts < minimum
    assert rows["Beta"]["dma_median"] > rows["Alpha"]["dma_median"]
    d = dma_client.get("/api/dma/correlation?group=acute&supplier=Alpha").json()
    assert d["n_trusts"] == 10  # 8 single + 2 running both; by-supplier card still shows the whole group
    assert len(d["by_supplier"]) == len(base["by_supplier"])
    assert dma_client.get("/api/dma/correlation?group=acute&category=PACS").json()["n_trusts"] == 20
    assert dma_client.get("/api/dma/correlation?supplier=Nobody").json()["n_trusts"] == 0
    t = next(x for x in base["trusts"] if x["ods_code"] == "T17")
    assert t["systems"]["EPR"] == ["Alpha", "Beta"]


def _itspend_book():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, mult in (("IT spend 2022-23", 1), ("IT spend 2023-24", 2), ("IT spend 2024-25", 3)):
        ws = wb.create_sheet(title)
        ws.append(["Trust ID", "Trust name", "Trust type", "Trust expenditure", "IT spend", "% of IT spend"])
        ws.append([7, "A Trust", "Acute Trust", 100_000_000, 1_000_000 * mult, 0.01 * mult])
        ws.append([99, "Unknown Trust", "Acute Trust", 50_000_000, 500_000, 0.01])
        ws.append([None, None, "Total", 150_000_000, 1_500_000, 0.01])
    h = wb.create_sheet("HIMSS")
    h.append(["Trust name", "Type", "HIMSS EMRAM", "HIMSS INFRAM", "Plan to Reassess?"])
    h.append(["A Trust", "Acute Trust", "5 (2022)", "Not Assessed", "No"])
    h.append(["Elsewhere", "Acute Trust", "N/A", "N/A", "N/A"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_itspend_import_joins_on_fhi_id_and_reports_unmapped(client, monkeypatch):
    from app import systems as sy
    monkeypatch.setattr(sy, "fhi_to_ods", lambda: {"7": "RXX"})
    r = client.post("/api/itspend/import", files={"file": ("it.xlsx", io.BytesIO(_itspend_book()), "application/octet-stream")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["organisations"] == 1 and body["unmapped_fhi_ids"] == {"99": "Unknown Trust"} and body["himss_values_loaded"] == 1
    from app.db import get_db
    from app.main import app
    db = next(app.dependency_overrides[get_db]()) if get_db in app.dependency_overrides else next(get_db())
    vals = {m.period: m.value for m in db.query(Metric).filter_by(ods_code="RXX", metric_code="itspend:it_spend")}
    assert vals == {"2022/23": 1_000_000, "2023/24": 2_000_000, "2024/25": 3_000_000}
    em = db.query(Metric).filter_by(ods_code="RXX", metric_code="itspend:himss_emram").one()
    assert em.value == 5 and em.quality_flag == "assessed 2022"
    assert db.query(Metric).filter_by(metric_code="itspend:it_spend_pct", period="2024/25").one().value == 3.0
    # re-upload replaces, never duplicates
    client.post("/api/itspend/import", files={"file": ("it.xlsx", io.BytesIO(_itspend_book()), "application/octet-stream")})
    assert db.query(Metric).filter_by(ods_code="RXX", metric_code="itspend:it_spend").count() == 3

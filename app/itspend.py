"""FHI IT-spend workbook (from FOIs sent to trusts) -> metrics, source_key "itspend".

PROPRIETARY: FHI's own survey data. Loaded through POST /api/itspend/import (or scripts/load_itspend.py), stored in the
database only, never committed.

Sheets used: "IT spend 2022-23", "IT spend 2023-24", "IT spend 2024-25" (Trust ID, total expenditure, IT spend, share) and
"HIMSS" (EMRAM and INFRAM stage, by trust name). Joined to ODS through the FHI Trust ID (trusts_ods.csv fhi_trust_id), the same
key as the installed-systems file. The derived sheets (percentage changes, rankings, analysis) are recomputed by the app, and the
EPR sheet is superseded by the master installations file; the DMA 2024 sheet duplicates the public NHS England release.
Trusts whose names changed through merger keep the FHI id, so earlier years sit under the successor's ODS code.
"""
import io
import re
from collections import Counter

import openpyxl

from app.periods import Period
from app.sources.loaders import LoadResult, _metric, to_float
from app import systems

SOURCE_KEY = "itspend"
YEARS = {"IT spend 2022-23": Period("2022/23", 202204, "year"), "IT spend 2023-24": Period("2023/24", 202304, "year"),
         "IT spend 2024-25": Period("2024/25", 202404, "year")}
SURVEY = Period("FHI survey Jan 2026", 202601, "month")  # file dated 26 Jan 2026; HIMSS stages carry their own year in the cell
_STAGE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*(?:\((\d{4})\))?\s*$")


def _norm(s) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower().replace("&", " and ")).strip()


def load_itspend(data: bytes, mapping: dict[str, str] | None = None) -> LoadResult:
    mapping = mapping if mapping is not None else systems.fhi_to_ods()
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = {ws.title.strip(): ws for ws in wb.worksheets}
    res = LoadResult(mapper_status="ok")
    unmapped, seen_name, per_year = {}, {}, Counter()
    flagged = Counter()
    for title, period in YEARS.items():
        ws = sheets.get(title)
        if ws is None:
            continue
        it = ws.iter_rows(values_only=True)
        hdr = [str(c or "").strip().lower() for c in next(it)]
        ix = {"id": hdr.index("trust id"), "name": hdr.index("trust name"),
              "exp": next(i for i, h in enumerate(hdr) if "expenditure" in h), "spend": hdr.index("it spend")}
        for r in it:
            tid = r[ix["id"]]
            if tid in (None, "") or not isinstance(tid, (int, float)):
                continue  # blank rows and the totals row
            fid, name = str(int(tid)), str(r[ix["name"]] or "").strip()
            ods = mapping.get(fid)
            if not ods:
                unmapped[fid] = name
                continue
            seen_name[fid] = name
            spend, exp = to_float(r[ix["spend"]]), to_float(r[ix["exp"]])
            res.orgs.setdefault(ods, (name, None))
            per_year[title] += 1
            if spend is not None:
                res.metrics.append(_metric("itspend:it_spend", "IT spend (FHI FOI survey)", ods, period, spend, SOURCE_KEY, "GBP"))
            if exp is not None:
                res.metrics.append(_metric("itspend:total_expenditure", "Total trust expenditure (as reported in FOI)", ods, period, exp, SOURCE_KEY, "GBP"))
            if spend is not None and exp:
                res.metrics.append(_metric("itspend:it_spend_pct", "IT spend as % of trust expenditure", ods, period, round(100 * spend / exp, 4), SOURCE_KEY, "%"))
    # HIMSS is by name: match to the names seen on the spend sheets (same trusts, same spelling), then to the ODS code
    by_name = {_norm(n): fid for fid, n in seen_name.items()}
    himss_unmatched, himss_rows = [], 0
    ws = sheets.get("HIMSS")
    if ws is not None:
        it = ws.iter_rows(values_only=True)
        hdr = [str(c or "").strip().lower() for c in next(it)]
        e, f = hdr.index("himss emram"), hdr.index("himss infram")
        for r in it:
            if not r[0]:
                continue
            fid = by_name.get(_norm(r[0]))
            ods = mapping.get(fid) if fid else None
            if not ods:
                if any(str(v).strip().upper() not in ("N/A", "NONE", "") for v in (r[e], r[f])):
                    himss_unmatched.append(str(r[0]).strip())  # only a problem if it carries a stage
                continue
            for col, code, label in ((e, "itspend:himss_emram", "HIMSS EMRAM stage (0-7)"), (f, "itspend:himss_infram", "HIMSS INFRAM stage (0-7)")):
                m = _STAGE.match(str(r[col] if r[col] is not None else ""))
                if not m:
                    continue
                row = _metric(code, label, ods, SURVEY, float(m.group(1)), SOURCE_KEY, "stage")
                if m.group(2):
                    row["quality_flag"] = "assessed " + m.group(2)
                    flagged["dated_stage"] += 1
                res.metrics.append(row)
                himss_rows += 1
    res.row_count = sum(per_year.values())
    res.profile = {
        "sheets_used": [t for t in list(YEARS) + ["HIMSS"] if t in sheets], "trust_rows_by_sheet": dict(per_year),
        "organisations": len(res.orgs), "unmapped_fhi_ids": unmapped, "himss_values_loaded": himss_rows,
        "himss_unmatched_with_a_stage": himss_unmatched, "himss_dated_stages": flagged["dated_stage"],
        "note": "Stages like '5 (2022)' load as 5 with quality_flag 'assessed 2022'. Text such as 'Planned for 2026/27' or 'Not Assessed' is skipped.",
    }
    return res


def store(db, file_name: str, raw: bytes, res: LoadResult) -> None:
    """Replace any earlier itspend load: one SourceFile, its metrics, and any orgs not yet known."""
    import hashlib

    from sqlalchemy import delete, insert, select

    from app.models import Metric, Org, SourceFile

    old = [i for (i,) in db.execute(select(SourceFile.id).where(SourceFile.source_key == SOURCE_KEY))]
    if old:
        db.execute(delete(Metric).where(Metric.source_file_id.in_(old)))
        db.execute(delete(SourceFile).where(SourceFile.id.in_(old)))
    sf = SourceFile(source_key=SOURCE_KEY, file_name=file_name, sha256=hashlib.sha256(raw).hexdigest(), size_bytes=len(raw),
                    status="loaded", mapper_status=res.mapper_status, row_count=res.row_count, metric_rows=len(res.metrics),
                    profile=res.profile, licence="FHI proprietary (FOI survey); not for redistribution")
    db.add(sf)
    for code, (name, kind) in res.orgs.items():
        if db.get(Org, code) is None:
            db.add(Org(ods_code=code, name=name, org_type=kind, first_seen_source=SOURCE_KEY))
    db.flush()
    rows = [dict(m, source_file_id=sf.id) for m in res.metrics]
    for i in range(0, len(rows), 2000):
        db.execute(insert(Metric), rows[i:i + 2000])
    db.commit()
